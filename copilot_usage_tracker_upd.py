#!/usr/bin/env python3
"""Track GitHub Copilot usage metrics for an enterprise, org, or team.

This script pulls Copilot metrics from GitHub's REST API and summarizes:
- Code acceptance metrics (suggestions vs acceptances, plus line acceptance)
- Model usage (aggregated across all model identifiers found in the payload)
- Chat ask details (total asks/chats and editor/channel breakdown)

Required permissions for the token depend on scope and org policy.
Typical usage requires a GitHub App or PAT with admin-level access.
"""

from __future__ import annotations

import argparse
import csv
import getpass
import glob
import json
import os
import re
import subprocess
import sys
import zipfile
from collections import defaultdict
from datetime import UTC, datetime, timedelta
from html import escape
from pathlib import Path
from typing import Any, Optional

import requests

try:
    from openpyxl import Workbook  # type: ignore[import-not-found]
except ImportError:  # pragma: no cover - optional dependency used only for Excel export
    Workbook = None

API_BASE = "https://api.github.com"
API_VERSION = "2022-11-28"

CONFIG_PATH = os.path.expanduser("~/.copilot_tracker.json")

# Default model rates (cost per token). Empty/zero means rate unknown.
# Populate with actual rates (cost per token) via --rates-file JSON if available.
DEFAULT_MODEL_RATES: dict[str, float] = {
    # example placeholders (tokens are individual tokens, not per-1k)
    "gpt-4": 0.0,
    "gpt-4o": 0.0,
    "gpt-4o-mini": 0.0,
    "gpt-4o-realtime": 0.0,
    "gpt-3.5-turbo": 0.0,
}


# ---------------------------------------------------------------------------
# Config file helpers
# ---------------------------------------------------------------------------

def load_config() -> dict[str, Any]:
    if not os.path.exists(CONFIG_PATH):
        return {}
    try:
        with open(CONFIG_PATH, "r", encoding="utf-8") as f:
            data = json.load(f)
        return data if isinstance(data, dict) else {}
    except (json.JSONDecodeError, OSError):
        return {}


def save_config(config: dict[str, Any]) -> None:
    try:
        with open(CONFIG_PATH, "w", encoding="utf-8") as f:
            json.dump(config, f, indent=2)
        os.chmod(CONFIG_PATH, 0o600)
    except OSError as exc:
        print(f"Warning: could not save config to {CONFIG_PATH}: {exc}", file=sys.stderr)


def ensure_parent_dir(path: str) -> None:
    Path(path).parent.mkdir(parents=True, exist_ok=True)


def _is_interactive() -> bool:
    return sys.stdin.isatty() and sys.stdout.isatty()


def _prompt(label: str, default: str | None = None, secret: bool = False) -> str:
    hint = f" [{default}]" if default else ""
    prompt_str = f"{label}{hint}: "
    if secret:
        value = getpass.getpass(prompt_str)
    else:
        value = input(prompt_str).strip()
    if not value and default:
        return default
    return value


def detect_github_org_from_git() -> str | None:
    """Try to detect GitHub org from git remote URL."""
    try:
        result = subprocess.run(
            ["git", "config", "--get", "remote.origin.url"],
            capture_output=True,
            text=True,
            timeout=5,
        )
        url = result.stdout.strip()
        if not url:
            return None
        
        # Parse github.com/owner/repo format
        if "github.com" in url:
            # Handle both https and ssh formats
            if url.startswith("git@"):
                # git@github.com:owner/repo.git
                parts = url.split("/")
                if len(parts) >= 2:
                    return parts[-2]
            elif url.startswith("https"):
                # https://github.com/owner/repo.git or .git
                parts = url.rstrip("/").split("/")
                if len(parts) >= 4:
                    return parts[-2]
        return None
    except Exception:
        return None


def detect_git_author(repo_path: str) -> str | None:
    for field in ("user.email", "user.name"):
        proc = subprocess.run(
            ["git", "-C", repo_path, "config", field],
            capture_output=True, text=True, check=False,
        )
        value = proc.stdout.strip()
        if value:
            return value
    return None


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Fetch and summarize GitHub Copilot usage metrics."
    )

    scope = parser.add_mutually_exclusive_group(required=False)
    scope.add_argument("--enterprise", help="GitHub Enterprise slug")
    scope.add_argument("--org", help="GitHub organization name")
    parser.add_argument(
        "--personal",
        action="store_true",
        help="Use personal mode (local git + optional ask logging)",
    )

    parser.add_argument(
        "--team",
        help="Team slug (only valid with --org); fetches team-level metrics",
    )
    parser.add_argument(
        "--repo-path",
        default=".",
        help="Repository path for personal mode git stats",
    )
    parser.add_argument(
        "--author",
        help="Optional git author filter for personal mode",
    )
    parser.add_argument(
        "--since",
        help="Start date (YYYY-MM-DD). Optional.",
    )
    parser.add_argument(
        "--until",
        help="End date (YYYY-MM-DD). Optional.",
    )
    parser.add_argument(
        "--token",
        default=os.getenv("GITHUB_TOKEN") or os.getenv("GH_TOKEN"),
        help="GitHub token (defaults to GITHUB_TOKEN/GH_TOKEN env var)",
    )
    parser.add_argument(
        "--raw-json",
        action="store_true",
        help="Print raw JSON payload after summary",
    )
    parser.add_argument(
        "--csv-out",
        help="Write daily summarized metrics to this CSV file path",
    )
    parser.add_argument(
        "--trend-out",
        help="Write an ASCII daily trend report to this text file path",
    )
    parser.add_argument(
        "--excel-out",
        help="Write daily rows and a summary sheet to this Excel (.xlsx) file path",
    )
    parser.add_argument(
        "--record-ask",
        help="Optional ask text to append into personal ask log",
    )
    parser.add_argument(
        "--ask-model",
        default="unknown",
        help="Model label used with --record-ask (example: gpt-4.1)",
    )
    parser.add_argument(
        "--ask-log-file",
        default="reports/personal_asks.jsonl",
        help="JSONL file path for personal ask logging",
    )
    parser.add_argument(
        "--no-vscode-sync",
        action="store_true",
        help="Skip auto-importing asks from VS Code Copilot Chat logs in personal mode",
    )
    parser.add_argument(
        "--vscode-log-dir",
        default=None,
        help="Override VS Code Copilot Chat log directory (default: auto-detect)",
    )
    parser.add_argument(
        "--save-config",
        action="store_true",
        help="Save current --org/--enterprise/--token/--author settings to ~/.copilot_tracker.json",
    )
    parser.add_argument(
        "--siebel-report",
        help="Write per-Siebel-developer license/tokens CSV file path",
    )
    parser.add_argument(
        "--license-summary-report",
        help="Write license-type totals (assigned/used) CSV file path",
    )
    parser.add_argument(
        "--rates-file",
        help="JSON file mapping model name to rate per token (float).",
    )
    parser.add_argument(
        "--cost-report",
        help="Write per-model token/cost CSV file path (sage cost)",
    )

    args = parser.parse_args()

    # --- Load saved config and apply as defaults for any unset values ---
    config = load_config()
    if not args.enterprise and not args.org and not args.personal:
        if config.get("enterprise"):
            args.enterprise = config["enterprise"]
        elif config.get("org"):
            args.org = config["org"]
    if not args.token:
        args.token = config.get("token") or os.getenv("GITHUB_TOKEN") or os.getenv("GH_TOKEN")
    if not args.author and config.get("author"):
        args.author = config["author"]

    if args.team and not args.org:
        parser.error("--team can only be used with --org")

    if args.personal and (args.enterprise or args.org or args.team):
        parser.error("--personal cannot be combined with --enterprise/--org/--team")

    # --- Interactive prompts for missing required fields ---
    if not args.personal and not (args.enterprise or args.org) and _is_interactive():
        print("No --org or --enterprise specified. Enter details to continue.")
        scope_choice = _prompt("Scope (org/enterprise)", default="org").lower()
        if scope_choice == "enterprise":
            args.enterprise = _prompt("Enterprise slug")
        else:
            args.org = _prompt("Organization name")

    if not args.personal and not (args.enterprise or args.org):
        parser.error("Use one of --enterprise/--org, or use --personal")

    if not args.token and not args.personal and _is_interactive():
        print("No GitHub token found. Enter a PAT with copilot metrics read access.")
        args.token = _prompt("GitHub token", secret=True)

    if not args.personal and not args.token:
        parser.error("Missing token: provide --token or set GITHUB_TOKEN/GH_TOKEN")

    for label, value in (("since", args.since), ("until", args.until)):
        if value:
            validate_date(value, f"--{label}")

    return args


def validate_date(value: str, arg_name: str) -> None:
    try:
        datetime.strptime(value, "%Y-%m-%d")
    except ValueError as exc:
        raise SystemExit(f"Invalid date for {arg_name}: {value}. Use YYYY-MM-DD") from exc


def build_endpoint(args: argparse.Namespace) -> str:
    if args.enterprise:
        return f"{API_BASE}/enterprises/{args.enterprise}/copilot/metrics"
    if args.team:
        return f"{API_BASE}/orgs/{args.org}/teams/{args.team}/copilot/metrics"
    return f"{API_BASE}/orgs/{args.org}/copilot/metrics"


def fetch_metrics(args: argparse.Namespace) -> list[dict[str, Any]]:
    url = build_endpoint(args)
    params: dict[str, str] = {}
    if args.since:
        params["since"] = args.since
    if args.until:
        params["until"] = args.until

    headers = {
        "Accept": "application/vnd.github+json",
        "Authorization": f"Bearer {args.token}",
        "X-GitHub-Api-Version": API_VERSION,
    }

    try:
        response = requests.get(url, headers=headers, params=params, timeout=60)
        response.raise_for_status()
    except requests.RequestException as exc:
        raise SystemExit(f"API request failed: {exc}") from exc

    try:
        data = response.json()
    except json.JSONDecodeError as exc:
        raise SystemExit("Unexpected API response: invalid JSON") from exc

    if not isinstance(data, list):
        raise SystemExit("Unexpected API response: expected a list of daily metrics")
    return data


def fetch_org_license_quota(args: argparse.Namespace) -> dict[str, float]:
    """Fetch organization license type token assignments for personal mode context.
    
    Returns a dict with license type as key and total assigned tokens as value.
    """
    if not (args.org or args.enterprise):
        return {}
    
    url = build_endpoint(args)
    params: dict[str, str] = {}
    # Get recent data to capture current quota
    if not args.since:
        # Last 7 days if not specified
        until_date = datetime.now().date()
        since_date = until_date - timedelta(days=7)
        params["since"] = str(since_date)
        params["until"] = str(until_date)
    else:
        params["since"] = args.since
        if args.until:
            params["until"] = args.until

    headers = {
        "Accept": "application/vnd.github+json",
        "Authorization": f"Bearer {args.token}",
        "X-GitHub-Api-Version": API_VERSION,
    }

    try:
        response = requests.get(url, headers=headers, params=params, timeout=60)
        response.raise_for_status()
    except requests.RequestException as exc:
        print(f"Warning: Could not fetch org metrics: {exc}")
        return {}

    try:
        data = response.json()
    except json.JSONDecodeError:
        print("Warning: Invalid response from org metrics endpoint")
        return {}

    if not isinstance(data, list):
        return {}

    # Aggregate by license type from the metrics
    siebel_devs = aggregate_siebel_developers(data)
    license_totals = aggregate_license_type_totals(siebel_devs)
    
    result = {}
    for license_type, stats in license_totals.items():
        assigned = int(round(stats.get("tokens_assigned", 0)))
        if assigned > 0:
            result[license_type] = assigned
    
    return result


def fetch_code_suggestion_kpis_for_scope(args: argparse.Namespace) -> dict[str, Any] | None:
    """Fetch code suggestion KPI counters from GitHub Copilot metrics endpoint."""
    if not (args.org or args.enterprise):
        return None
    if not getattr(args, "token", None):
        return None

    try:
        metrics = fetch_metrics(args)
    except SystemExit:
        return None

    return summarize_code_suggestion_kpis(metrics)


def iter_dicts(payload: Any):
    if isinstance(payload, dict):
        yield payload
        for value in payload.values():
            yield from iter_dicts(value)
    elif isinstance(payload, list):
        for item in payload:
            yield from iter_dicts(item)


def sum_keys(payload: Any, key_names: set[str]) -> float:
    total = 0.0
    for node in iter_dicts(payload):
        for key, value in node.items():
            if key in key_names and isinstance(value, (int, float)):
                total += float(value)
    return total


def aggregate_models(metrics: list[dict[str, Any]]) -> dict[str, dict[str, float]]:
    model_stats: dict[str, dict[str, float]] = defaultdict(lambda: defaultdict(float))

    for node in iter_dicts(metrics):
        model_name = node.get("model") or node.get("model_name")
        if not isinstance(model_name, str) or not model_name.strip():
            continue

        name = model_name.strip()
        for key, value in node.items():
            if key in {"model", "model_name"}:
                continue
            if isinstance(value, (int, float)):
                model_stats[name][key] += float(value)

    return model_stats


def aggregate_chat_breakdown(metrics: list[dict[str, Any]]) -> dict[str, float]:
    breakdown: dict[str, float] = defaultdict(float)
    chat_keys = {
        "total_chats",
        "total_chat_requests",
        "total_chat_turns",
        "total_asks",
        "total_messages",
    }

    for node in iter_dicts(metrics):
        for key in chat_keys:
            value = node.get(key)
            if isinstance(value, (int, float)):
                # Keep an at-a-glance metric by key.
                breakdown[key] += float(value)

    return dict(breakdown)


def aggregate_siebel_developers(metrics: list[dict[str, Any]]) -> dict[str, dict[str, float | str]]:
    """Aggregate per-developer license type, tokens assigned and tokens used.

    The function is tolerant to varying payload field names and will try several
    common key names to locate email/identifier and numeric token fields.
    """
    devs: dict[str, dict[str, float | str]] = defaultdict(lambda: {
        "license_type": "unknown",
        "tokens_assigned": 0.0,
        "tokens_used": 0.0,
    })

    assigned_keys = ("tokens_assigned", "total_tokens_assigned", "assigned_tokens")
    used_keys = ("tokens_used", "tokens_consumed", "usage_tokens")
    license_keys = ("license_type", "license", "license_name")
    id_keys = ("developer_email", "user_email", "email", "author_email", "login", "developer")

    for node in iter_dicts(metrics):
        # try to identify a developer identifier (prefer email/login)
        dev_id: Optional[str] = None
        for k in id_keys:
            v = node.get(k)
            if isinstance(v, str) and v.strip():
                dev_id = v.strip()
                break

        if not dev_id:
            continue

        # license
        for k in license_keys:
            v = node.get(k)
            if isinstance(v, str) and v.strip():
                devs[dev_id]["license_type"] = v.strip()

        # numeric fields
        for k in assigned_keys:
            v = node.get(k)
            if isinstance(v, (int, float)):
                devs[dev_id]["tokens_assigned"] += float(v)
        for k in used_keys:
            v = node.get(k)
            if isinstance(v, (int, float)):
                devs[dev_id]["tokens_used"] += float(v)

    return dict(devs)


def aggregate_license_type_totals(
    siebel_devs: dict[str, dict[str, float | str]],
) -> dict[str, dict[str, float]]:
    """Aggregate token assignment and usage totals grouped by license type."""
    totals: dict[str, dict[str, float]] = defaultdict(
        lambda: {"tokens_assigned": 0.0, "tokens_used": 0.0}
    )

    for stats in siebel_devs.values():
        license_type = stats.get("license_type", "unknown")
        if not isinstance(license_type, str) or not license_type.strip():
            license_type = "unknown"
        name = license_type.strip()

        assigned = stats.get("tokens_assigned", 0)
        used = stats.get("tokens_used", 0)

        if isinstance(assigned, (int, float)):
            totals[name]["tokens_assigned"] += float(assigned)
        if isinstance(used, (int, float)):
            totals[name]["tokens_used"] += float(used)

    return dict(totals)


def load_rate_map(path: str | None) -> dict[str, float]:
    rates = dict(DEFAULT_MODEL_RATES)
    if not path:
        return rates
    try:
        with open(path, "r", encoding="utf-8") as fh:
            data = json.load(fh)
        if isinstance(data, dict):
            for k, v in data.items():
                try:
                    rates[str(k)] = float(v)
                except Exception:
                    continue
    except OSError:
        pass
    return rates


def aggregate_model_token_usage(metrics: list[dict[str, Any]]) -> dict[str, dict[str, float]]:
    """Aggregate token usage per model from the metrics payload.

    This helper uses exact token counts when present, while also accepting
    a broad set of common token field names and estimating totals from
    input/output/cached counts if needed.
    """
    usage: dict[str, dict[str, float]] = defaultdict(lambda: {
        "input_tokens": 0.0,
        "output_tokens": 0.0,
        "cached_tokens": 0.0,
        "total_tokens": 0.0,
    })

    input_keys = (
        "total_input_tokens",
        "input_tokens",
        "tokens_in",
        "prompt_tokens",
        "prompt_token_count",
        "request_tokens",
    )
    output_keys = (
        "total_output_tokens",
        "output_tokens",
        "tokens_out",
        "completion_tokens",
        "completion_token_count",
        "response_tokens",
    )
    cached_keys = ("total_cached_tokens", "cached_tokens")
    total_keys = (
        "total_tokens",
        "tokens_total",
        "token_count",
        "total_token_count",
    )

    for node in iter_dicts(metrics):
        model = node.get("model") or node.get("model_name")
        if not isinstance(model, str) or not model.strip():
            continue
        name = model.strip()

        for k in input_keys:
            v = node.get(k)
            if isinstance(v, (int, float)):
                usage[name]["input_tokens"] += float(v)
        for k in output_keys:
            v = node.get(k)
            if isinstance(v, (int, float)):
                usage[name]["output_tokens"] += float(v)
        for k in cached_keys:
            v = node.get(k)
            if isinstance(v, (int, float)):
                usage[name]["cached_tokens"] += float(v)
        for k in total_keys:
            v = node.get(k)
            if isinstance(v, (int, float)):
                usage[name]["total_tokens"] += float(v)

    # compute totals if not explicitly reported
    for name, stats in usage.items():
        if not stats.get("total_tokens"):
            stats["total_tokens"] = (
                stats.get("input_tokens", 0.0)
                + stats.get("output_tokens", 0.0)
                + stats.get("cached_tokens", 0.0)
            )

    return dict(usage)


def estimate_personal_tokens(ask_events: list[dict[str, Any]], rate_map: dict[str, float] | None = None) -> dict[str, Any]:
    """Estimate token usage from personal ask events using a simple heuristic.

    Uses `ask_length` (characters) if present, otherwise falls back to len(ask).
    Assumes ~4 characters per token to estimate input tokens. Returns totals
    and per-model estimates and optional cost when a rate_map is provided.
    """
    total_chars = 0
    total_asks = 0
    total_input_tokens = 0
    total_output_tokens = 0
    total_cost = 0.0
    per_model_chars: dict[str, int] = defaultdict(int)
    per_model_tokens: dict[str, int] = defaultdict(int)

    for ev in ask_events:
        length = None
        if isinstance(ev.get("ask_length"), int):
            length = ev.get("ask_length")
        else:
            a = ev.get("ask")
            if isinstance(a, str):
                length = len(a)

        if not length:
            continue

        total_chars += int(length)
        total_asks += 1
        explicit_input = ev.get("estimated_input_tokens")
        if isinstance(explicit_input, int):
            total_input_tokens += explicit_input
        else:
            total_input_tokens += _estimate_text_tokens(str(ev.get("ask") or ""))

        explicit_output = ev.get("estimated_output_tokens")
        if isinstance(explicit_output, int):
            total_output_tokens += explicit_output

        explicit_cost = ev.get("estimated_total_cost")
        if isinstance(explicit_cost, (int, float)):
            total_cost += float(explicit_cost)

        model = ev.get("model") or ev.get("ask_model") or "unknown"
        if not isinstance(model, str) or not model:
            model = "unknown"
        model = _normalize_model_name(model)
        per_model_chars[model] += int(length)
        if isinstance(explicit_input, int):
            per_model_tokens[model] += explicit_input

    # Fallback to heuristic when enriched tokens are not present.
    if not per_model_tokens:
        for m, chars in per_model_chars.items():
            per_model_tokens[m] = int(round(chars / 4.0))

    total_tokens = total_input_tokens + total_output_tokens
    if total_tokens == 0:
        total_tokens = sum(per_model_tokens.values())

    if rate_map:
        for model_name, tokens in per_model_tokens.items():
            try:
                rate = float(rate_map.get(model_name, rate_map.get(model_name.lower(), 0.0)))
            except Exception:
                rate = 0.0
            total_cost += tokens * rate

    if not per_model_chars and total_chars == 0:
        return {
            "total_asks": total_asks,
            "total_chars": total_chars,
            "input_tokens": 0,
            "output_tokens": 0,
            "total_tokens": 0,
            "per_model_tokens": {},
            "estimated_cost": 0.0,
        }

    return {
        "total_asks": total_asks,
        "total_chars": total_chars,
        "input_tokens": total_input_tokens,
        "output_tokens": total_output_tokens,
        "total_tokens": total_tokens,
        "per_model_tokens": dict(per_model_tokens),
        "estimated_cost": round(total_cost, 6),
    }


def summarize_personal_billing(ask_events: list[dict[str, Any]]) -> dict[str, Any]:
    sku_counts: dict[str, int] = defaultdict(int)
    priced_events = 0
    input_tokens = 0
    output_tokens = 0
    total_tokens = 0
    total_cost = 0.0
    models_with_billing: dict[str, dict[str, float]] = {}

    for event in ask_events:
        sku = event.get("copilot_sku")
        if isinstance(sku, str) and sku.strip():
            sku_counts[sku.strip()] += 1

        model = event.get("normalized_model") or event.get("model") or "unknown"
        if not isinstance(model, str) or not model.strip():
            model = "unknown"
        model_name = model.strip()

        event_input_tokens = event.get("estimated_input_tokens")
        event_output_tokens = event.get("estimated_output_tokens")
        event_total_tokens = event.get("estimated_total_tokens")

        if isinstance(event_input_tokens, int):
            input_tokens += event_input_tokens
        if isinstance(event_output_tokens, int):
            output_tokens += event_output_tokens
        if isinstance(event_total_tokens, int):
            total_tokens += event_total_tokens
        elif isinstance(event_input_tokens, int) or isinstance(event_output_tokens, int):
            total_tokens += int(event_input_tokens or 0) + int(event_output_tokens or 0)

        event_cost = event.get("estimated_total_cost")
        if isinstance(event_cost, (int, float)):
            total_cost += float(event_cost)
            priced_events += 1

        billing_input = event.get("billing_input_price")
        billing_output = event.get("billing_output_price")
        billing_cache = event.get("billing_cache_price")
        billing_batch_size = event.get("billing_batch_size")
        if any(isinstance(v, (int, float)) for v in (billing_input, billing_output, billing_cache)):
            models_with_billing[model_name] = {
                "billing_input_price": float(billing_input or 0.0),
                "billing_output_price": float(billing_output or 0.0),
                "billing_cache_price": float(billing_cache or 0.0),
                "billing_batch_size": float(billing_batch_size or 1_000_000.0),
            }

    return {
        "copilot_skus": dict(sku_counts),
        "priced_events": priced_events,
        "input_tokens": input_tokens,
        "output_tokens": output_tokens,
        "total_tokens": total_tokens,
        "estimated_cost": round(total_cost, 6),
        "models_with_billing": models_with_billing,
    }


def _license_type_from_sku(sku_value: str) -> str:
    sku = sku_value.strip().lower()
    if "enterprise" in sku:
        return "enterprise"
    if "business" in sku:
        return "business"
    if "pro_plus" in sku or "pro-plus" in sku:
        return "pro_plus"
    if "pro" in sku:
        return "pro"
    if "trial" in sku:
        return "individual_trial"
    return sku_value.strip() or "unknown"


def derive_license_types_from_skus(sku_counts: dict[str, int]) -> list[str]:
    types = {_license_type_from_sku(sku) for sku in sku_counts.keys() if sku}
    if not types:
        return ["unknown"]
    return sorted(types)


def summarize_code_suggestion_kpis(payload: Any) -> dict[str, Any]:
    """Summarize code suggestion KPIs from any payload shape.

    Works for org metrics payloads and returns N/A markers when fields
    are not present (common in personal-mode local ask logs).
    """
    shown = sum_keys(payload, {"total_code_suggestions"})
    accepted = sum_keys(payload, {"total_code_acceptances"})
    lines_accepted = sum_keys(payload, {"total_code_lines_accepted"})

    shown_int = int(round(shown))
    accepted_int = int(round(accepted))
    lines_accepted_int = int(round(lines_accepted))
    acceptance_pct = (lines_accepted / shown * 100) if shown > 0 else 0.0
    available = shown_int > 0 or accepted_int > 0 or lines_accepted_int > 0

    return {
        "shown": shown_int,
        "accepted": accepted_int,
        "lines_of_code_accepted": lines_accepted_int,
        "acceptance_pct": acceptance_pct,
        "available": available,
    }


def extract_personal_token_quota(ask_events: list[dict[str, Any]]) -> dict[str, Any]:
    """Extract total assigned/used token quota values from personal ask events."""
    assigned_keys = {
        "tokens_assigned",
        "total_tokens_assigned",
        "assigned_tokens",
    }
    used_keys = {
        "tokens_used",
        "total_tokens_used",
        "used_tokens",
    }

    total_assigned = 0
    total_used = 0
    for ev in ask_events:
        for node in iter_dicts(ev):
            for key, value in node.items():
                if key in assigned_keys and isinstance(value, (int, float)):
                    total_assigned += int(round(value))
                if key in used_keys and isinstance(value, (int, float)):
                    total_used += int(round(value))

    remaining = max(0, total_assigned - total_used)
    usage_pct = (total_used / total_assigned * 100) if total_assigned else 0.0
    return {
        "total_assigned_tokens": total_assigned,
        "total_used_tokens": total_used,
        "total_remaining_tokens": remaining,
        "usage_percent": usage_pct,
    }


def find_vscode_transcript_files() -> list[str]:
    files: list[str] = []
    search_roots = [
        os.path.expanduser("~/.vscode-remote/data/User/workspaceStorage"),
        os.path.expanduser("~/.config/Code/User/workspaceStorage"),
        os.path.expanduser("~/Library/Application Support/Code/User/workspaceStorage"),
        os.path.expanduser("~/AppData/Roaming/Code/User/workspaceStorage"),
    ]

    for root in search_roots:
        if not os.path.isdir(root):
            continue
        pattern = os.path.join(root, "*", "GitHub.copilot-chat", "transcripts", "*.jsonl")
        files.extend(glob.glob(pattern))

    return sorted(set(files))


def parse_vscode_transcript(
    transcript_path: str | None = None,
    since: str | None = None,
    until: str | None = None,
) -> dict[str, Any]:
    """Parse VS Code Copilot Chat transcript to extract messages and estimate tokens.

    Reads the transcript JSONL file and extracts user (input) and assistant (output) messages.
    Estimates tokens using ~4 characters per token heuristic.

    Returns dict with:
        - input_chars, output_chars, total_chars
        - input_tokens, output_tokens, total_tokens (estimated)
        - user_messages, assistant_messages (counts)
    """
    transcript_files = [transcript_path] if transcript_path else find_vscode_transcript_files()

    try:
        if not transcript_files:
            return {
                "input_chars": 0,
                "output_chars": 0,
                "total_chars": 0,
                "input_tokens": 0,
                "output_tokens": 0,
                "total_tokens": 0,
                "user_messages": 0,
                "assistant_messages": 0,
            }

        user_messages = []
        assistant_messages = []

        for raw_path in transcript_files:
            transcript_file = Path(raw_path)
            if not transcript_file.exists():
                continue

            with open(transcript_file, "r", encoding="utf-8") as fh:
                for line in fh:
                    try:
                        obj = json.loads(line)
                    except Exception:
                        continue

                    timestamp = obj.get("timestamp")
                    if isinstance(timestamp, str) and len(timestamp) >= 10:
                        date_value = timestamp[:10]
                        if since and date_value < since:
                            continue
                        if until and date_value > until:
                            continue

                    event_type = obj.get("type", "")
                    data = obj.get("data", {})

                    if event_type == "user.message":
                        content = data.get("content", "")
                        if isinstance(content, str) and content:
                            user_messages.append(content)

                    elif event_type == "assistant.message":
                        content = data.get("content", "")
                        if isinstance(content, str) and content:
                            assistant_messages.append(content)

        # Calculate totals
        input_chars = sum(len(m) for m in user_messages)
        output_chars = sum(len(m) for m in assistant_messages)
        total_chars = input_chars + output_chars

        # Estimate tokens: ~4 chars per token
        input_tokens = int(round(input_chars / 4.0))
        output_tokens = int(round(output_chars / 4.0))
        total_tokens = input_tokens + output_tokens

        return {
            "input_chars": input_chars,
            "output_chars": output_chars,
            "total_chars": total_chars,
            "input_tokens": input_tokens,
            "output_tokens": output_tokens,
            "total_tokens": total_tokens,
            "user_messages": len(user_messages),
            "assistant_messages": len(assistant_messages),
        }
    except Exception as e:
        # Return zeroes on any error
        return {
            "input_chars": 0,
            "output_chars": 0,
            "total_chars": 0,
            "input_tokens": 0,
            "output_tokens": 0,
            "total_tokens": 0,
            "user_messages": 0,
            "assistant_messages": 0,
        }


def build_daily_rows(metrics: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []

    for day in metrics:
        date_value = day.get("date")
        if not isinstance(date_value, str):
            date_value = "unknown"

        suggestions = sum_keys(day, {"total_code_suggestions"})
        acceptances = sum_keys(day, {"total_code_acceptances"})
        lines_suggested = sum_keys(day, {"total_code_lines_suggested"})
        lines_accepted = sum_keys(day, {"total_code_lines_accepted"})
        asks = sum_keys(
            day,
            {
                "total_chats",
                "total_chat_requests",
                "total_chat_turns",
                "total_asks",
                "total_messages",
            },
        )
        active_users = sum_keys(day, {"total_active_users"})
        engaged_users = sum_keys(day, {"total_engaged_users"})

        acceptance_rate = (acceptances / suggestions) * 100 if suggestions else 0.0
        line_acceptance_rate = (
            (lines_accepted / lines_suggested) * 100 if lines_suggested else 0.0
        )

        rows.append(
            {
                "date": date_value,
                "code_suggestions": int(round(suggestions)),
                "code_acceptances": int(round(acceptances)),
                "acceptance_rate_pct": round(acceptance_rate, 2),
                "lines_suggested": int(round(lines_suggested)),
                "lines_accepted": int(round(lines_accepted)),
                "line_acceptance_rate_pct": round(line_acceptance_rate, 2),
                "chat_ask_signals": int(round(asks)),
                "active_users": int(round(active_users)),
                "engaged_users": int(round(engaged_users)),
            }
        )

    rows.sort(key=lambda row: (_parse_date(row["date"]), row["date"]))
    return rows


def write_daily_csv(rows: list[dict[str, Any]], output_path: str) -> None:
    ensure_parent_dir(output_path)

    fieldnames = [
        "date",
        "code_suggestions",
        "code_acceptances",
        "acceptance_rate_pct",
        "lines_suggested",
        "lines_accepted",
        "line_acceptance_rate_pct",
        "chat_ask_signals",
        "active_users",
        "engaged_users",
    ]

    with open(output_path, "w", encoding="utf-8", newline="") as csv_file:
        writer = csv.DictWriter(csv_file, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def _bar(value: float, max_value: float, width: int = 24) -> str:
    if max_value <= 0:
        return ""
    filled = int(round((value / max_value) * width))
    if filled < 0:
        filled = 0
    if filled > width:
        filled = width
    return "#" * filled


def _parse_date(date_value: str) -> datetime:
    try:
        return datetime.strptime(date_value, "%Y-%m-%d")
    except ValueError:
        return datetime.max


def build_trend_report(rows: list[dict[str, Any]]) -> str:
    if not rows:
        return "No rows available for trend report."

    max_acceptance = max(float(r["acceptance_rate_pct"]) for r in rows)
    max_asks = max(float(r["chat_ask_signals"]) for r in rows)
    max_suggestions = max(float(r["code_suggestions"]) for r in rows)

    lines = []
    lines.append("GitHub Copilot Daily Trend Report")
    lines.append("=" * 72)
    lines.append("A) Acceptance Rate (%)")
    for row in rows:
        value = float(row["acceptance_rate_pct"])
        lines.append(
            f"{row['date']} | {_bar(value, max_acceptance):<24} {value:6.2f}%"
        )

    lines.append("")
    lines.append("B) Chat Ask Signals")
    for row in rows:
        value = float(row["chat_ask_signals"])
        lines.append(f"{row['date']} | {_bar(value, max_asks):<24} {value:8.0f}")

    lines.append("")
    lines.append("C) Code Suggestions")
    for row in rows:
        value = float(row["code_suggestions"])
        lines.append(
            f"{row['date']} | {_bar(value, max_suggestions):<24} {value:8.0f}"
        )

    return "\n".join(lines)


def write_trend_report(rows: list[dict[str, Any]], output_path: str) -> None:
    ensure_parent_dir(output_path)

    report = build_trend_report(rows)
    with open(output_path, "w", encoding="utf-8") as text_file:
        text_file.write(report + "\n")


def write_siebel_csv(devs: dict[str, dict[str, float | str]], output_path: str) -> None:
    ensure_parent_dir(output_path)
    fieldnames = ["developer", "license_type", "tokens_assigned", "tokens_used"]
    with open(output_path, "w", encoding="utf-8", newline="") as csv_file:
        writer = csv.DictWriter(csv_file, fieldnames=fieldnames)
        writer.writeheader()
        for dev, stats in sorted(devs.items()):
            writer.writerow(
                {
                    "developer": dev,
                    "license_type": stats.get("license_type", "unknown"),
                    "tokens_assigned": int(round(stats.get("tokens_assigned", 0)))
                    if isinstance(stats.get("tokens_assigned", 0), (int, float))
                    else stats.get("tokens_assigned", 0),
                    "tokens_used": int(round(stats.get("tokens_used", 0)))
                    if isinstance(stats.get("tokens_used", 0), (int, float))
                    else stats.get("tokens_used", 0),
                }
            )


def write_license_summary_csv(
    license_totals: dict[str, dict[str, float]],
    output_path: str,
) -> None:
    ensure_parent_dir(output_path)
    fieldnames = ["license_type", "tokens_assigned", "tokens_used"]

    with open(output_path, "w", encoding="utf-8", newline="") as csv_file:
        writer = csv.DictWriter(csv_file, fieldnames=fieldnames)
        writer.writeheader()
        for license_type, stats in sorted(license_totals.items()):
            writer.writerow(
                {
                    "license_type": license_type,
                    "tokens_assigned": int(round(stats.get("tokens_assigned", 0))),
                    "tokens_used": int(round(stats.get("tokens_used", 0))),
                }
            )


def write_cost_csv(costs: dict[str, dict[str, float]], output_path: str, rate_map: dict[str, float] | None = None) -> None:
    ensure_parent_dir(output_path)
    fieldnames = [
        "model",
        "input_tokens",
        "output_tokens",
        "cached_tokens",
        "total_tokens",
        "rate_per_token",
        "sage_cost",
    ]
    with open(output_path, "w", encoding="utf-8", newline="") as csv_file:
        writer = csv.DictWriter(csv_file, fieldnames=fieldnames)
        writer.writeheader()
        for model, stats in sorted(costs.items()):
            rate = 0.0
            if rate_map and model in rate_map:
                rate = rate_map[model]
            total = float(stats.get("total_tokens", 0.0))
            cost = total * float(rate)
            writer.writerow(
                {
                    "model": model,
                    "input_tokens": int(round(stats.get("input_tokens", 0))),
                    "output_tokens": int(round(stats.get("output_tokens", 0))),
                    "cached_tokens": int(round(stats.get("cached_tokens", 0))),
                    "total_tokens": int(round(total)),
                    "rate_per_token": rate,
                    "sage_cost": round(cost, 6),
                }
            )


def write_org_excel(
    rows: list[dict[str, Any]],
    siebel_devs: dict[str, dict[str, float | str]],
    metrics: list[dict[str, Any]],
    rate_map: dict[str, float] | None,
    output_path: str,
) -> None:
    ensure_parent_dir(output_path)

    daily_headers = [
        "date",
        "code_suggestions",
        "code_acceptances",
        "acceptance_rate_pct",
        "lines_suggested",
        "lines_accepted",
        "line_acceptance_rate_pct",
        "chat_ask_signals",
        "active_users",
        "engaged_users",
    ]

    siebel_rows = [["developer", "license_type", "tokens_assigned", "tokens_used"]]
    for dev, stats in sorted(siebel_devs.items()):
        siebel_rows.append([
            dev,
            stats.get("license_type", "unknown"),
            int(round(stats.get("tokens_assigned", 0))) if isinstance(stats.get("tokens_assigned", 0), (int, float)) else stats.get("tokens_assigned", 0),
            int(round(stats.get("tokens_used", 0))) if isinstance(stats.get("tokens_used", 0), (int, float)) else stats.get("tokens_used", 0),
        ])

    license_totals = aggregate_license_type_totals(siebel_devs)
    license_types = sorted(license_totals.keys()) if license_totals else ["unknown"]
    license_rows = [["license_type", "tokens_assigned", "tokens_used"]]
    for license_type, stats in sorted(license_totals.items()):
        license_rows.append([
            license_type,
            int(round(stats.get("tokens_assigned", 0))),
            int(round(stats.get("tokens_used", 0))),
        ])

    daily_rows = [daily_headers] + [[r[h] for h in daily_headers] for r in rows]

    total_shown = sum(int(r.get("code_suggestions", 0)) for r in rows)
    total_accepted = sum(int(r.get("code_acceptances", 0)) for r in rows)
    total_lines_accepted = sum(int(r.get("lines_accepted", 0)) for r in rows)
    total_lines_suggested = sum(int(r.get("lines_suggested", 0)) for r in rows)
    total_chat_signals = sum(int(r.get("chat_ask_signals", 0)) for r in rows)
    acceptance_pct = (total_lines_accepted / total_shown * 100) if total_shown else 0.0

    token_usage = aggregate_model_token_usage(metrics)
    grand_total_tokens = 0
    grand_total_cost = 0.0
    token_model_rows = [[
        "model",
        "input_tokens",
        "output_tokens",
        "cached_tokens",
        "total_tokens",
        "rate_per_token",
        "estimated_cost",
    ]]
    for model_name, stats in sorted(token_usage.items()):
        inp = int(round(stats.get("input_tokens", 0)))
        out = int(round(stats.get("output_tokens", 0)))
        cached = int(round(stats.get("cached_tokens", 0)))
        total = float(stats.get("total_tokens", 0.0))
        grand_total_tokens += int(round(total))
        rate = 0.0
        if rate_map:
            rate = rate_map.get(model_name, rate_map.get(model_name.lower(), 0.0))
        cost = total * float(rate)
        grand_total_cost += cost
        token_model_rows.append([
            model_name,
            inp,
            out,
            cached,
            int(round(total)),
            rate,
            round(cost, 6),
        ])

    token_kpi_rows = [
        ["Metric", "Value"],
        ["License type(s)", ", ".join(license_types)],
        ["Shown", total_shown],
        ["Accepted", total_accepted],
        ["Lines of code accepted", total_lines_accepted],
        ["% Acceptance", round(acceptance_pct, 2)],
        ["Lines suggested", total_lines_suggested],
        ["Chat ask signals", total_chat_signals],
        ["Total Copilot tokens", grand_total_tokens],
        ["Estimated token cost", round(grand_total_cost, 6)],
    ]

    if Workbook is not None:
        wb = Workbook()
        ws_daily = wb.active
        ws_daily.title = "Daily"
        for row in daily_rows:
            ws_daily.append(row)

        ws_siebel = wb.create_sheet("Siebel Developers")
        for row in siebel_rows:
            ws_siebel.append(row)

        ws_license = wb.create_sheet("License Summary")
        for row in license_rows:
            ws_license.append(row)

        ws_token_kpi = wb.create_sheet("Token and KPI")
        for row in token_kpi_rows:
            ws_token_kpi.append(row)

        ws_token_models = wb.create_sheet("Token by Model")
        for row in token_model_rows:
            ws_token_models.append(row)

        wb.save(output_path)
        return

    _write_xlsx_fallback(
        output_path,
        [
            ("Daily", daily_rows),
            ("Siebel Developers", siebel_rows),
            ("License Summary", license_rows),
            ("Token and KPI", token_kpi_rows),
            ("Token by Model", token_model_rows),
        ],
    )


def resolve_personal_window(since: str | None, until: str | None) -> tuple[str, str]:
    today = datetime.now(UTC).date()

    if until:
        end_date = datetime.strptime(until, "%Y-%m-%d").date()
    else:
        end_date = today

    if since:
        start_date = datetime.strptime(since, "%Y-%m-%d").date()
    else:
        start_date = end_date - timedelta(days=29)

    if start_date > end_date:
        raise SystemExit("Invalid range: --since must be <= --until")

    return start_date.isoformat(), end_date.isoformat()


def _date_span(start_date: str, end_date: str) -> list[str]:
    start = datetime.strptime(start_date, "%Y-%m-%d").date()
    end = datetime.strptime(end_date, "%Y-%m-%d").date()
    dates: list[str] = []
    current = start
    while current <= end:
        dates.append(current.isoformat())
        current += timedelta(days=1)
    return dates


def get_git_daily_stats(
    repo_path: str,
    since: str,
    until: str,
    author: str | None,
) -> dict[str, dict[str, int]]:
    cmd = [
        "git",
        "-C",
        repo_path,
        "log",
        f"--since={since}",
        f"--until={until} 23:59:59",
        "--date=short",
        "--pretty=format:__COMMIT__%H|%ad",
        "--numstat",
    ]
    if author:
        cmd.append(f"--author={author}")

    proc = subprocess.run(cmd, capture_output=True, text=True, check=False)
    if proc.returncode != 0:
        raise SystemExit(f"git log failed: {proc.stderr.strip()}")

    stats: dict[str, dict[str, int]] = defaultdict(
        lambda: {"commits": 0, "files_changed": 0, "lines_added": 0, "lines_deleted": 0}
    )

    current_date: str | None = None
    for raw_line in proc.stdout.splitlines():
        line = raw_line.strip()
        if not line:
            continue

        if line.startswith("__COMMIT__"):
            parts = line.split("|", 1)
            if len(parts) == 2:
                current_date = parts[1]
                if current_date:
                    stats[current_date]["commits"] += 1
            continue

        if current_date is None:
            continue

        fields = raw_line.split("\t")
        if len(fields) < 3:
            continue

        added, deleted = fields[0], fields[1]
        stats[current_date]["files_changed"] += 1
        if added.isdigit():
            stats[current_date]["lines_added"] += int(added)
        if deleted.isdigit():
            stats[current_date]["lines_deleted"] += int(deleted)

    return dict(stats)


# ---------------------------------------------------------------------------
# VS Code Copilot Chat log auto-sync
# ---------------------------------------------------------------------------

_MODEL_NAME_SANITIZE_RE = re.compile(r"-[0-9]{4}-[0-9]{2}-[0-9]{2}$")

_VSCODE_LOG_SEARCH_PATHS = [
    os.path.expanduser("~/.vscode-remote/data/logs"),
    os.path.expanduser("~/.config/Code/logs"),
    os.path.expanduser("~/Library/Application Support/Code/logs"),
    os.path.expanduser("~/AppData/Roaming/Code/logs"),
]

_CCREQ_SUCCESS_RE = re.compile(
    r'^(\d{4}-\d{2}-\d{2}) \d{2}:\d{2}:\d{2}\.\d+\s+\[info\]\s+'
    r'(ccreq:[0-9a-fA-F]+)\.copilotmd \| success \| ([^|]+?) \| \d+ms \| \[([^\]]+)\]'
)
_CCREQ_SUCCESS_WITH_DURATION_RE = re.compile(
    r'^(\d{4}-\d{2}-\d{2}) \d{2}:\d{2}:\d{2}\.\d+\s+\[info\]\s+'
    r'(ccreq:[0-9a-fA-F]+)\.copilotmd \| success \| ([^|]+?) \| (\d+)ms \| \[([^\]]+)\]'
)
_COPILOT_SKU_RE = re.compile(
    r'^(\d{4}-\d{2}-\d{2}).*?copilot token sku:\s+([^\s]+)'
)


def _estimate_text_tokens(text: str) -> int:
    return int(round(len(text) / 4.0)) if text else 0


def _normalize_model_name(model_name: str) -> str:
    name = model_name.strip()
    if not name:
        return "unknown"
    if " -> " in name:
        name = name.split(" -> ", 1)[0].strip()
    return _MODEL_NAME_SANITIZE_RE.sub("", name)


def _build_model_name_candidates(model_name: str) -> list[str]:
    normalized = _normalize_model_name(model_name)
    candidates = [normalized]
    lowered = normalized.lower()
    if lowered not in candidates:
        candidates.append(lowered)
    without_family_date = _MODEL_NAME_SANITIZE_RE.sub("", lowered)
    if without_family_date not in candidates:
        candidates.append(without_family_date)
    return [candidate for candidate in candidates if candidate]


def _extract_default_billing(model_entry: dict[str, Any]) -> dict[str, float] | None:
    billing = model_entry.get("billing")
    if not isinstance(billing, dict):
        return None

    token_prices = billing.get("token_prices")
    if not isinstance(token_prices, dict):
        return None

    default_prices = token_prices.get("default")
    if not isinstance(default_prices, dict):
        return None

    batch_size = token_prices.get("batch_size", 1_000_000)
    if not isinstance(batch_size, (int, float)) or batch_size <= 0:
        batch_size = 1_000_000

    result: dict[str, float] = {"billing_batch_size": float(batch_size)}
    for source_key, target_key in (
        ("input_price", "billing_input_price"),
        ("output_price", "billing_output_price"),
        ("cache_price", "billing_cache_price"),
        ("context_max", "billing_context_max_tokens"),
    ):
        value = default_prices.get(source_key)
        if isinstance(value, (int, float)):
            result[target_key] = float(value)

    restricted_to = billing.get("restricted_to")
    if isinstance(restricted_to, list):
        result["billing_restricted_to_count"] = float(len(restricted_to))

    return result if len(result) > 1 else None


def find_vscode_model_catalog_files() -> list[str]:
    files: list[str] = []
    search_roots = [
        os.path.expanduser("~/.vscode-remote/data/User/workspaceStorage"),
        os.path.expanduser("~/.config/Code/User/workspaceStorage"),
        os.path.expanduser("~/Library/Application Support/Code/User/workspaceStorage"),
        os.path.expanduser("~/AppData/Roaming/Code/User/workspaceStorage"),
    ]

    for root in search_roots:
        if not os.path.isdir(root):
            continue
        pattern = os.path.join(root, "*", "GitHub.copilot-chat", "debug-logs", "*", "models.json")
        files.extend(glob.glob(pattern))

    return sorted(set(files))


def load_vscode_model_billing_map() -> dict[str, dict[str, float]]:
    billing_map: dict[str, dict[str, float]] = {}

    for path in find_vscode_model_catalog_files():
        try:
            with open(path, "r", encoding="utf-8") as fh:
                payload = json.load(fh)
        except (OSError, json.JSONDecodeError):
            continue

        if not isinstance(payload, list):
            continue

        for entry in payload:
            if not isinstance(entry, dict):
                continue
            model_id = entry.get("id") or entry.get("version") or entry.get("name")
            if not isinstance(model_id, str) or not model_id.strip():
                continue

            billing = _extract_default_billing(entry)
            if not billing:
                continue

            for candidate in _build_model_name_candidates(model_id):
                if candidate not in billing_map:
                    billing_map[candidate] = dict(billing)

    return billing_map


def detect_copilot_sku(log_files: list[str], since: str, until: str) -> str | None:
    sku_value: str | None = None
    for log_path in log_files:
        try:
            with open(log_path, "r", encoding="utf-8", errors="replace") as fh:
                for raw in fh:
                    match = _COPILOT_SKU_RE.match(raw)
                    if not match:
                        continue
                    date_str, parsed_sku = match.group(1), match.group(2).strip()
                    if date_str < since or date_str > until:
                        continue
                    if parsed_sku:
                        sku_value = parsed_sku
        except OSError:
            continue
    return sku_value


def enrich_personal_ask_event(
    entry: dict[str, Any],
    model_billing_map: dict[str, dict[str, float]] | None = None,
    copilot_sku: str | None = None,
) -> dict[str, Any]:
    event = dict(entry)

    # Ensure code-suggestion KPI keys exist in personal logs, even when upstream
    # local VS Code success logs do not provide these counters.
    for key in (
        "total_code_suggestions",
        "total_code_acceptances",
        "total_code_lines_accepted",
    ):
        value = event.get(key)
        if isinstance(value, (int, float)):
            event[key] = int(round(value))
        else:
            event[key] = 0

    ask_text = event.get("ask") if isinstance(event.get("ask"), str) else ""
    ask_length = event.get("ask_length")
    if not isinstance(ask_length, int):
        ask_length = len(ask_text)
    event["ask_length"] = ask_length

    estimated_input_tokens = event.get("estimated_input_tokens")
    if not isinstance(estimated_input_tokens, int):
        estimated_input_tokens = _estimate_text_tokens(ask_text)
    event["estimated_input_tokens"] = estimated_input_tokens

    estimated_output_tokens = event.get("estimated_output_tokens")
    if not isinstance(estimated_output_tokens, int):
        estimated_output_tokens = 0
    event["estimated_output_tokens"] = estimated_output_tokens

    estimated_total_tokens = event.get("estimated_total_tokens")
    if not isinstance(estimated_total_tokens, int):
        estimated_total_tokens = estimated_input_tokens + estimated_output_tokens
    event["estimated_total_tokens"] = estimated_total_tokens

    model_name = event.get("model") if isinstance(event.get("model"), str) else "unknown"
    normalized_model = _normalize_model_name(model_name)
    event["normalized_model"] = normalized_model
    event.setdefault("token_details_source", "estimated")

    if copilot_sku and not event.get("copilot_sku"):
        event["copilot_sku"] = copilot_sku

    if model_billing_map:
        billing: dict[str, float] | None = None
        for candidate in _build_model_name_candidates(model_name):
            billing = model_billing_map.get(candidate)
            if billing:
                break

        if billing:
            for key, value in billing.items():
                event[key] = value

            batch_size = float(billing.get("billing_batch_size", 1_000_000.0) or 1_000_000.0)
            input_price = float(billing.get("billing_input_price", 0.0))
            output_price = float(billing.get("billing_output_price", 0.0))
            event["estimated_input_cost"] = round((estimated_input_tokens * input_price) / batch_size, 6)
            event["estimated_output_cost"] = round((estimated_output_tokens * output_price) / batch_size, 6)
            event["estimated_total_cost"] = round(
                float(event["estimated_input_cost"]) + float(event["estimated_output_cost"]),
                6,
            )

    return event


def count_vscode_success_lines(
    log_files: list[str],
    since: str,
    until: str,
) -> int:
    """Count success ask log lines in range for capture-rate visibility."""
    total = 0
    for log_path in log_files:
        try:
            with open(log_path, "r", encoding="utf-8", errors="replace") as fh:
                for raw in fh:
                    if "copilotmd | success |" not in raw:
                        continue
                    if len(raw) < 10:
                        continue
                    date_str = raw[:10]
                    if not re.match(r"^\d{4}-\d{2}-\d{2}$", date_str):
                        continue
                    if date_str < since or date_str > until:
                        continue
                    total += 1
        except OSError:
            continue
    return total


def find_vscode_copilot_log_files(override_dir: str | None = None) -> list[str]:
    search_dirs = [override_dir] if override_dir else _VSCODE_LOG_SEARCH_PATHS
    found: list[str] = []
    for base in search_dirs:
        if not base or not os.path.isdir(base):
            continue
        pattern = os.path.join(base, "*", "exthost1", "GitHub.copilot-chat", "GitHub Copilot Chat.log")
        found.extend(glob.glob(pattern))
    return sorted(set(found))


def parse_vscode_log_asks(
    log_files: list[str],
    since: str,
    until: str,
    model_billing_map: dict[str, dict[str, float]] | None = None,
    copilot_sku: str | None = None,
) -> list[dict[str, Any]]:
    events: list[dict[str, Any]] = []
    for log_path in log_files:
        try:
            with open(log_path, "r", encoding="utf-8", errors="replace") as fh:
                for raw in fh:
                    m = _CCREQ_SUCCESS_WITH_DURATION_RE.match(raw)
                    if not m:
                        continue
                    date_str, req_id, model_raw, channel_raw = (
                        m.group(1),
                        m.group(2),
                        m.group(3).strip(),
                        m.group(5).strip(),
                    )
                    duration_ms = int(m.group(4))
                    if date_str < since or date_str > until:
                        continue
                    model_name = _normalize_model_name(model_raw)
                    channel = channel_raw.split("/", 1)[0].strip() if channel_raw else "unknown"
                    # Reconstruct a UTC timestamp string from log line (date only)
                    events.append(enrich_personal_ask_event({
                        "timestamp_utc": date_str + "T00:00:00Z",
                        "date": date_str,
                        "model": model_name,
                        "channel": channel,
                        "ask_length": 0,
                        "ask": "",
                        "source": "vscode_log",
                        "request_id": req_id,
                        "duration_ms": duration_ms,
                    }, model_billing_map=model_billing_map, copilot_sku=copilot_sku))
        except OSError:
            continue
    return events


def sync_vscode_asks(
    ask_log_file: str,
    since: str,
    until: str,
    override_dir: str | None,
) -> dict[str, int]:
    """Import asks from VS Code logs and return sync/capture statistics."""
    log_files = find_vscode_copilot_log_files(override_dir)
    model_billing_map = load_vscode_model_billing_map()
    copilot_sku = detect_copilot_sku(log_files, since, until)
    if not log_files:
        return {
            "log_files": 0,
            "success_lines": 0,
            "parsed_success": 0,
            "new_added": 0,
            "billing_models": len(model_billing_map),
        }

    success_lines = count_vscode_success_lines(log_files, since, until)

    existing_events = load_personal_ask_events(ask_log_file)
    seen_ids: set[str] = {
        e["request_id"] for e in existing_events if "request_id" in e
    }

    new_events = parse_vscode_log_asks(
        log_files,
        since,
        until,
        model_billing_map=model_billing_map,
        copilot_sku=copilot_sku,
    )
    to_add = [e for e in new_events if e["request_id"] not in seen_ids]

    if not to_add:
        return {
            "log_files": len(log_files),
            "success_lines": success_lines,
            "parsed_success": len(new_events),
            "new_added": 0,
            "billing_models": len(model_billing_map),
        }

    ensure_parent_dir(ask_log_file)

    with open(ask_log_file, "a", encoding="utf-8") as logf:
        for entry in to_add:
            logf.write(json.dumps(entry, ensure_ascii=True) + "\n")

    return {
        "log_files": len(log_files),
        "success_lines": success_lines,
        "parsed_success": len(new_events),
        "new_added": len(to_add),
        "billing_models": len(model_billing_map),
    }


def append_personal_ask_event(
    ask_log_file: str,
    ask_text: str,
    model: str,
) -> None:
    ensure_parent_dir(ask_log_file)

    now = datetime.now(UTC)
    entry = enrich_personal_ask_event({
        "timestamp_utc": now.strftime("%Y-%m-%dT%H:%M:%SZ"),
        "date": now.date().isoformat(),
        "model": model,
        "ask_length": len(ask_text),
        "ask": ask_text,
        "source": "manual",
    }, model_billing_map=load_vscode_model_billing_map())

    with open(ask_log_file, "a", encoding="utf-8") as logf:
        logf.write(json.dumps(entry, ensure_ascii=True) + "\n")


def load_personal_ask_events(ask_log_file: str) -> list[dict[str, Any]]:
    if not os.path.exists(ask_log_file):
        return []

    events: list[dict[str, Any]] = []
    model_billing_map = load_vscode_model_billing_map()
    with open(ask_log_file, "r", encoding="utf-8") as logf:
        for line in logf:
            line = line.strip()
            if not line:
                continue
            try:
                payload = json.loads(line)
            except json.JSONDecodeError:
                continue
            if isinstance(payload, dict):
                events.append(enrich_personal_ask_event(payload, model_billing_map=model_billing_map))
    return events


def filter_personal_ask_events(
    ask_events: list[dict[str, Any]],
    since: str,
    until: str,
) -> list[dict[str, Any]]:
    filtered: list[dict[str, Any]] = []
    for event in ask_events:
        date_value = event.get("date")
        if not isinstance(date_value, str):
            continue
        if since <= date_value <= until:
            filtered.append(event)
    return filtered


def build_personal_daily_rows(
    git_stats: dict[str, dict[str, int]],
    ask_events: list[dict[str, Any]],
    since: str,
    until: str,
) -> list[dict[str, Any]]:
    ask_daily_counts: dict[str, int] = defaultdict(int)
    ask_daily_chars: dict[str, int] = defaultdict(int)

    for event in ask_events:
        date_value = event.get("date")
        if not isinstance(date_value, str):
            continue
        ask_daily_counts[date_value] += 1

        ask_length = event.get("ask_length")
        if isinstance(ask_length, int):
            ask_daily_chars[date_value] += ask_length

    rows: list[dict[str, Any]] = []
    for date_value in _date_span(since, until):
        day_git = git_stats.get(date_value, {})
        added = int(day_git.get("lines_added", 0))
        deleted = int(day_git.get("lines_deleted", 0))
        rows.append(
            {
                "date": date_value,
                "commits": int(day_git.get("commits", 0)),
                "files_changed": int(day_git.get("files_changed", 0)),
                "lines_added": added,
                "lines_deleted": deleted,
                "net_lines": added - deleted,
                "ask_count": int(ask_daily_counts.get(date_value, 0)),
                "ask_characters": int(ask_daily_chars.get(date_value, 0)),
            }
        )

    return rows


def summarize_ask_models(ask_events: list[dict[str, Any]]) -> dict[str, int]:
    counts: dict[str, int] = defaultdict(int)
    for event in ask_events:
        model = event.get("model")
        if isinstance(model, str) and model.strip():
            counts[model.strip()] += 1
        else:
            counts["unknown"] += 1
    return dict(counts)


def print_personal_summary(
    rows: list[dict[str, Any]],
    ask_events: list[dict[str, Any]],
    org_license_quota: dict[str, float] | None = None,
    current_copilot_sku: str | None = None,
    code_kpi_override: dict[str, Any] | None = None,
) -> None:
    if not rows:
        print("No personal rows generated for the selected date range.")
        return

    summary = build_personal_summary(
        rows,
        ask_events,
        code_kpi_override=code_kpi_override,
    )

    print("=" * 72)
    print("Personal Tracking Summary")
    print("=" * 72)
    print(f"Days covered      : {summary['days_covered']}")
    print(f"Date range        : {summary['date_range']}")
    print("\nGit Activity")
    print("-" * 72)
    print(f"Commits           : {summary['commits']:,}")
    print(f"Files changed     : {summary['files_changed']:,}")
    print(f"Lines added       : {summary['lines_added']:,}")
    print(f"Lines deleted     : {summary['lines_deleted']:,}")
    print(f"Net lines         : {summary['net_lines']:,}")

    print("\nAsk Activity")
    print("-" * 72)
    print(f"Asks logged       : {summary['asks_logged']:,}")
    print(f"Ask characters    : {summary['ask_characters']:,}")

    model_counts = summary["ask_model_breakdown"]
    if model_counts:
        print("Ask model breakdown:")
        for model, count in sorted(model_counts.items()):
            print(f"  {model}: {count:,}")

    kpis = summary.get("code_suggestion_kpis", {})
    print("\nCode Suggestion KPIs")
    print("-" * 72)
    if kpis.get("available"):
        print(
            f"Shown                           : {int(kpis.get('shown', 0)):,} "
            "(Number of GitHub code suggestions provided to developers)"
        )
        print(
            f"Accepted                        : {int(kpis.get('accepted', 0)):,} "
            "(Number of suggestions actually used/copied)"
        )
        print(
            f"Lines of code accepted          : {int(kpis.get('lines_of_code_accepted', 0)):,} "
            "(Lines of code inserted into the editor from Copilot suggestions)"
        )
        print(
            f"% Acceptance                    : {float(kpis.get('acceptance_pct', 0.0)):.2f}% "
            "(Lines of Code accepted / shown * 100)"
        )
    else:
        print("Shown                           : N/A")
        print("Accepted                        : N/A")
        print("Lines of code accepted          : N/A")
        print("% Acceptance                    : N/A")
        print("Code suggestion counters are not available in personal-mode local logs.")

    billing_summary = summary.get("billing_summary", {})
    license_types = summary.get("license_types", ["unknown"])
    print("\nLicensing and Billing")
    print("-" * 72)
    print(f"License type(s)               : {', '.join(license_types)}")
    sku_counts = dict(billing_summary.get("copilot_skus", {}))
    if current_copilot_sku and current_copilot_sku not in sku_counts:
        sku_counts[current_copilot_sku] = 0
    if sku_counts:
        print("Copilot SKU(s):")
        for sku_name, count in sorted(sku_counts.items()):
            suffix = f" ({count:,} logged ask events)" if count else ""
            print(f"  {sku_name}{suffix}")
    else:
        print("Copilot SKU       : unavailable from local logs")

    print(f"Ask input tokens  : {int(billing_summary.get('input_tokens', 0)):,}")
    print(f"Ask output tokens : {int(billing_summary.get('output_tokens', 0)):,}")
    print(f"Ask total tokens  : {int(billing_summary.get('total_tokens', 0)):,}")
    print(f"Priced ask events : {int(billing_summary.get('priced_events', 0)):,}")
    print(f"Estimated ask cost: {float(billing_summary.get('estimated_cost', 0.0)):,.6f}")
    if (
        int(billing_summary.get("total_tokens", 0)) == 0
        and int(summary.get("transcript_tokens", {}).get("total_tokens", 0)) > 0
    ):
        print("Per-request token counts are not present in success-line imports; see transcript totals below.")

    models_with_billing = billing_summary.get("models_with_billing", {})
    if isinstance(models_with_billing, dict) and models_with_billing:
        print("Model billing rates (per batch):")
        for model_name, pricing in sorted(models_with_billing.items()):
            batch_size = int(round(float(pricing.get("billing_batch_size", 1_000_000.0))))
            input_price = float(pricing.get("billing_input_price", 0.0))
            output_price = float(pricing.get("billing_output_price", 0.0))
            cache_price = float(pricing.get("billing_cache_price", 0.0))
            print(
                f"  {model_name}: input={input_price:g}, output={output_price:g}, cache={cache_price:g} per {batch_size:,} tokens"
            )

    # Display transcript token usage
    transcript_data = summary.get("transcript_tokens", {})
    if transcript_data.get("total_tokens", 0) > 0:
        print("\nTranscript Token Usage (from VS Code Chat)")
        print("-" * 72)
        print(f"Total input tokens (user messages)  : {transcript_data['input_tokens']:,}")
        print(f"Total output tokens (assistant msg) : {transcript_data['output_tokens']:,}")
        print(f"Total tokens consumed              : {transcript_data['total_tokens']:,}")
        print(f"Messages: {transcript_data['user_messages']} user, {transcript_data['assistant_messages']} assistant")

    # Display cumulative token usage
    total_est_tokens = summary.get("total_estimated_tokens", 0)
    token_quota = summary.get("personal_token_quota", {})
    quota_assigned = int(token_quota.get("total_assigned_tokens", 0))
    quota_used = int(token_quota.get("total_used_tokens", 0))
    quota_remaining = int(token_quota.get("total_remaining_tokens", 0))
    quota_pct = token_quota.get("usage_percent", 0.0)

    if total_est_tokens > 0:
        print("\n" + "=" * 72)
        print(f"Total token usage (estimated)       : {total_est_tokens:,}")
        if transcript_data.get("total_tokens", 0) > 0:
            print(f"Total tokens consumed (transcript)  : {transcript_data['total_tokens']:,}")
    else:
        print("\n" + "=" * 72)
        print("Total token usage: not available")

    if quota_assigned > 0 or quota_used > 0:
        print("\nQuota Token Balance")
        print("-" * 72)
        print(f"Total tokens assigned : {quota_assigned:,}")
        print(f"Total tokens used     : {quota_used:,}")
        print(f"Total tokens remaining: {quota_remaining:,}")
        print(f"Token usage percent   : {quota_pct:.2f}%")
    else:
        print("\nQuota Token Balance")
        print("-" * 72)
        print("No personal quota fields found in imported ask events.")
        print("Total tokens assigned remains unavailable for this dataset.")

    # Display organization license type quota if available
    if org_license_quota:
        print("\nOrganization License Type Token Quota")
        print("-" * 72)
        for license_type in sorted(org_license_quota.keys()):
            assigned = int(org_license_quota[license_type])
            print(f"{license_type:<24}: {assigned:,} tokens assigned")


def build_personal_summary(
    rows: list[dict[str, Any]],
    ask_events: list[dict[str, Any]],
    code_kpi_override: dict[str, Any] | None = None,
) -> dict[str, Any]:
    total_commits = sum(r["commits"] for r in rows)
    total_files_changed = sum(r["files_changed"] for r in rows)
    total_added = sum(r["lines_added"] for r in rows)
    total_deleted = sum(r["lines_deleted"] for r in rows)
    total_asks = sum(r["ask_count"] for r in rows)
    total_ask_chars = sum(r["ask_characters"] for r in rows)

    date_range = "unknown"
    if rows:
        date_range = f"{rows[0]['date']} to {rows[-1]['date']}"

    # Parse transcript for token data
    since = rows[0]["date"] if rows else None
    until = rows[-1]["date"] if rows else None
    transcript_data = parse_vscode_transcript(since=since, until=until)
    
    # Prefer ask-log totals when available, otherwise fall back to transcript totals.
    token_quota = extract_personal_token_quota(ask_events)
    billing_summary = summarize_personal_billing(ask_events)
    license_types = derive_license_types_from_skus(
        dict(billing_summary.get("copilot_skus", {}))
    )
    code_suggestion_kpis = code_kpi_override or summarize_code_suggestion_kpis(ask_events)
    total_estimated_tokens = int(billing_summary.get("total_tokens", 0))
    if total_estimated_tokens <= 0:
        total_estimated_tokens = int(transcript_data.get("total_tokens", 0))

    return {
        "days_covered": len(rows),
        "date_range": date_range,
        "commits": total_commits,
        "files_changed": total_files_changed,
        "lines_added": total_added,
        "lines_deleted": total_deleted,
        "net_lines": total_added - total_deleted,
        "asks_logged": total_asks,
        "ask_characters": total_ask_chars,
        "ask_model_breakdown": summarize_ask_models(ask_events),
        "transcript_tokens": transcript_data,
        "total_estimated_tokens": total_estimated_tokens,
        "total_transcript_tokens": transcript_data.get("total_tokens", 0),
        "personal_token_quota": token_quota,
        "billing_summary": billing_summary,
        "license_types": license_types,
        "code_suggestion_kpis": code_suggestion_kpis,
    }


def write_personal_csv(rows: list[dict[str, Any]], output_path: str) -> None:
    ensure_parent_dir(output_path)

    fieldnames = [
        "date",
        "commits",
        "files_changed",
        "lines_added",
        "lines_deleted",
        "net_lines",
        "ask_count",
        "ask_characters",
    ]

    with open(output_path, "w", encoding="utf-8", newline="") as csv_file:
        writer = csv.DictWriter(csv_file, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def build_personal_trend_report(rows: list[dict[str, Any]]) -> str:
    if not rows:
        return "No rows available for trend report."

    max_commits = max(float(r["commits"]) for r in rows)
    max_asks = max(float(r["ask_count"]) for r in rows)
    max_changed = max(float(r["files_changed"]) for r in rows)

    lines: list[str] = []
    lines.append("Personal Daily Trend Report")
    lines.append("=" * 72)
    lines.append("A) Commits")
    for row in rows:
        value = float(row["commits"])
        lines.append(f"{row['date']} | {_bar(value, max_commits):<24} {value:8.0f}")

    lines.append("")
    lines.append("B) Ask Count")
    for row in rows:
        value = float(row["ask_count"])
        lines.append(f"{row['date']} | {_bar(value, max_asks):<24} {value:8.0f}")

    lines.append("")
    lines.append("C) Files Changed")
    for row in rows:
        value = float(row["files_changed"])
        lines.append(f"{row['date']} | {_bar(value, max_changed):<24} {value:8.0f}")

    return "\n".join(lines)


def write_personal_trend_report(rows: list[dict[str, Any]], output_path: str) -> None:
    ensure_parent_dir(output_path)

    report = build_personal_trend_report(rows)
    with open(output_path, "w", encoding="utf-8") as text_file:
        text_file.write(report + "\n")


def write_personal_excel(
    rows: list[dict[str, Any]],
    ask_events: list[dict[str, Any]],
    output_path: str,
) -> None:
    ensure_parent_dir(output_path)

    headers = [
        "date",
        "commits",
        "files_changed",
        "lines_added",
        "lines_deleted",
        "net_lines",
        "ask_count",
        "ask_characters",
    ]

    summary = build_personal_summary(rows, ask_events)
    transcript = summary.get("transcript_tokens", {})
    billing = summary.get("billing_summary", {})
    license_types = summary.get("license_types", ["unknown"])
    kpis = summary.get("code_suggestion_kpis", {})
    quota = summary.get("personal_token_quota", {})

    # Prefer per-ask token counters when present; otherwise use transcript totals
    # so Excel aligns with what the personal console summary displays.
    ask_input_tokens = int(billing.get("input_tokens", 0))
    ask_output_tokens = int(billing.get("output_tokens", 0))
    ask_total_tokens = int(billing.get("total_tokens", 0))
    token_source = "ask_events"
    if ask_total_tokens <= 0:
        ask_input_tokens = int(transcript.get("input_tokens", 0))
        ask_output_tokens = int(transcript.get("output_tokens", 0))
        ask_total_tokens = int(transcript.get("total_tokens", 0))
        token_source = "transcript_estimated"

    daily_rows = [headers] + [[row[h] for h in headers] for row in rows]
    summary_rows = [
        ["Metric", "Value"],
        ["Days covered", summary["days_covered"]],
        ["Date range", summary["date_range"]],
        ["License type(s)", ", ".join(license_types)],
        ["Commits", summary["commits"]],
        ["Files changed", summary["files_changed"]],
        ["Lines added", summary["lines_added"]],
        ["Lines deleted", summary["lines_deleted"]],
        ["Net lines", summary["net_lines"]],
        ["Asks logged", summary["asks_logged"]],
        ["Ask characters", summary["ask_characters"]],
        ["Total token usage (estimated)", summary.get("total_estimated_tokens", 0)],
        ["Transcript input tokens", transcript.get("input_tokens", 0)],
        ["Transcript output tokens", transcript.get("output_tokens", 0)],
        ["Transcript total tokens", transcript.get("total_tokens", 0)],
    ]
    model_rows = [["Model", "Count"]] + [
        [model, count] for model, count in sorted(summary["ask_model_breakdown"].items())
    ]

    token_kpi_rows = [
        ["Metric", "Value"],
        ["License type(s)", ", ".join(license_types)],
        ["Token source", token_source],
        ["Shown", int(kpis.get("shown", 0)) if kpis.get("available") else "N/A"],
        ["Accepted", int(kpis.get("accepted", 0)) if kpis.get("available") else "N/A"],
        [
            "Lines of code accepted",
            int(kpis.get("lines_of_code_accepted", 0)) if kpis.get("available") else "N/A",
        ],
        [
            "% Acceptance",
            round(float(kpis.get("acceptance_pct", 0.0)), 2) if kpis.get("available") else "N/A",
        ],
        ["Ask input tokens", ask_input_tokens],
        ["Ask output tokens", ask_output_tokens],
        ["Ask total tokens", ask_total_tokens],
        ["Estimated ask cost", float(billing.get("estimated_cost", 0.0))],
        ["Total tokens assigned", int(quota.get("total_assigned_tokens", 0))],
        ["Total tokens used", int(quota.get("total_used_tokens", 0))],
        ["Total tokens remaining", int(quota.get("total_remaining_tokens", 0))],
        ["Token usage percent", round(float(quota.get("usage_percent", 0.0)), 2)],
    ]

    if Workbook is not None:
        wb = Workbook()
        ws_daily = wb.active
        ws_daily.title = "Daily"
        for row in daily_rows:
            ws_daily.append(row)

        ws_summary = wb.create_sheet("Summary")
        for row in summary_rows:
            ws_summary.append(row)

        ws_models = wb.create_sheet("Ask Model Breakdown")
        for row in model_rows:
            ws_models.append(row)

        ws_token_kpi = wb.create_sheet("Token and KPI")
        for row in token_kpi_rows:
            ws_token_kpi.append(row)

        wb.save(output_path)
        return

    _write_xlsx_fallback(
        output_path,
        [
            ("Daily", daily_rows),
            ("Summary", summary_rows),
            ("Ask Model Breakdown", model_rows),
            ("Token and KPI", token_kpi_rows),
        ],
    )



def _write_xlsx_fallback(output_path: str, sheets: list[tuple[str, list[list[Any]]]]) -> None:
    def cell_xml(value: Any) -> str:
        if value is None:
            return "<c/>"
        if isinstance(value, bool):
            return f"<c t=\"b\"><v>{int(value)}</v></c>"
        if isinstance(value, (int, float)):
            return f"<c><v>{value}</v></c>"
        return f"<c t=\"inlineStr\"><is><t>{escape(str(value))}</t></is></c>"

    def sheet_xml(rows: list[list[Any]]) -> str:
        row_nodes: list[str] = []
        for idx, row in enumerate(rows, start=1):
            cells = "".join(cell_xml(v) for v in row)
            row_nodes.append(f"<row r=\"{idx}\">{cells}</row>")
        rows_xml = "".join(row_nodes)
        return (
            "<?xml version=\"1.0\" encoding=\"UTF-8\" standalone=\"yes\"?>"
            "<worksheet xmlns=\"http://schemas.openxmlformats.org/spreadsheetml/2006/main\">"
            f"<sheetData>{rows_xml}</sheetData>"
            "</worksheet>"
        )

    content_types_overrides = []
    rel_nodes = []
    workbook_sheet_nodes = []
    sheet_blobs: list[tuple[str, str]] = []

    for idx, (name, rows) in enumerate(sheets, start=1):
        escaped_name = escape(name)
        workbook_sheet_nodes.append(
            f"<sheet name=\"{escaped_name}\" sheetId=\"{idx}\" r:id=\"rId{idx}\"/>"
        )
        rel_nodes.append(
            f"<Relationship Id=\"rId{idx}\" "
            "Type=\"http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet\" "
            f"Target=\"worksheets/sheet{idx}.xml\"/>"
        )
        content_types_overrides.append(
            f"<Override PartName=\"/xl/worksheets/sheet{idx}.xml\" "
            "ContentType=\"application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml\"/>"
        )
        sheet_blobs.append((f"xl/worksheets/sheet{idx}.xml", sheet_xml(rows)))

    workbook_xml = (
        "<?xml version=\"1.0\" encoding=\"UTF-8\" standalone=\"yes\"?>"
        "<workbook xmlns=\"http://schemas.openxmlformats.org/spreadsheetml/2006/main\" "
        "xmlns:r=\"http://schemas.openxmlformats.org/officeDocument/2006/relationships\">"
        f"<sheets>{''.join(workbook_sheet_nodes)}</sheets>"
        "</workbook>"
    )

    workbook_rels_xml = (
        "<?xml version=\"1.0\" encoding=\"UTF-8\" standalone=\"yes\"?>"
        "<Relationships xmlns=\"http://schemas.openxmlformats.org/package/2006/relationships\">"
        f"{''.join(rel_nodes)}"
        "</Relationships>"
    )

    root_rels_xml = (
        "<?xml version=\"1.0\" encoding=\"UTF-8\" standalone=\"yes\"?>"
        "<Relationships xmlns=\"http://schemas.openxmlformats.org/package/2006/relationships\">"
        "<Relationship Id=\"rId1\" "
        "Type=\"http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument\" "
        "Target=\"xl/workbook.xml\"/>"
        "</Relationships>"
    )

    content_types_xml = (
        "<?xml version=\"1.0\" encoding=\"UTF-8\" standalone=\"yes\"?>"
        "<Types xmlns=\"http://schemas.openxmlformats.org/package/2006/content-types\">"
        "<Default Extension=\"rels\" ContentType=\"application/vnd.openxmlformats-package.relationships+xml\"/>"
        "<Default Extension=\"xml\" ContentType=\"application/xml\"/>"
        "<Override PartName=\"/xl/workbook.xml\" "
        "ContentType=\"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml\"/>"
        f"{''.join(content_types_overrides)}"
        "</Types>"
    )

    with zipfile.ZipFile(output_path, mode="w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", content_types_xml)
        zf.writestr("_rels/.rels", root_rels_xml)
        zf.writestr("xl/workbook.xml", workbook_xml)
        zf.writestr("xl/_rels/workbook.xml.rels", workbook_rels_xml)
        for path, blob in sheet_blobs:
            zf.writestr(path, blob)


def print_summary(metrics: list[dict[str, Any]], rate_map: dict[str, float] | None = None) -> None:
    if not metrics:
        print("No Copilot metrics returned for the selected scope/date range.")
        return

    dates = [m.get("date") for m in metrics if isinstance(m.get("date"), str)]
    date_min = min(dates) if dates else "unknown"
    date_max = max(dates) if dates else "unknown"

    total_suggestions = sum_keys(metrics, {"total_code_suggestions"})
    total_acceptances = sum_keys(metrics, {"total_code_acceptances"})
    total_lines_suggested = sum_keys(metrics, {"total_code_lines_suggested"})
    total_lines_accepted = sum_keys(metrics, {"total_code_lines_accepted"})
    efficiency_acceptance_pct = (
        (total_lines_accepted / total_suggestions) * 100 if total_suggestions else 0.0
    )

    acceptance_rate = (
        (total_acceptances / total_suggestions) * 100 if total_suggestions else 0.0
    )
    line_acceptance_rate = (
        (total_lines_accepted / total_lines_suggested) * 100
        if total_lines_suggested
        else 0.0
    )

    total_active_users = sum_keys(metrics, {"total_active_users"})
    total_engaged_users = sum_keys(metrics, {"total_engaged_users"})

    chat_breakdown = aggregate_chat_breakdown(metrics)
    model_usage = aggregate_models(metrics)
    token_usage = aggregate_model_token_usage(metrics)
    siebel_devs = aggregate_siebel_developers(metrics)

    print("=" * 72)
    print("GitHub Copilot Usage Summary")
    print("=" * 72)
    print(f"Days returned: {len(metrics)}")
    print(f"Date range : {date_min} to {date_max}")
    print("\nCode Acceptance")
    print("-" * 72)
    print(f"Code suggestions : {total_suggestions:,.0f}")
    print(f"Code acceptances : {total_acceptances:,.0f}")
    print(f"Acceptance rate  : {acceptance_rate:.2f}%")
    print(f"Lines suggested  : {total_lines_suggested:,.0f}")
    print(f"Lines accepted   : {total_lines_accepted:,.0f}")
    print(f"Line accept rate : {line_acceptance_rate:.2f}%")

    # Requested KPI view using business labels.
    print("\nRequested KPI View")
    print("-" * 72)
    print(
        f"Shown                           : {total_suggestions:,.0f} "
        "(Number of GitHub code suggestions provided to developers)"
    )
    print(
        f"Accepted                        : {total_acceptances:,.0f} "
        "(Number of suggestions actually used/copied)"
    )
    print(
        f"Lines of code accepted          : {total_lines_accepted:,.0f} "
        "(Lines inserted into the editor from Copilot suggestions)"
    )
    print(
        f"% Acceptance                    : {efficiency_acceptance_pct:.2f}% "
        "(Efficiency metric = Lines of Code accepted / shown)"
    )

    print("\nEngagement")
    print("-" * 72)
    print(f"Active users     : {total_active_users:,.0f}")
    print(f"Engaged users    : {total_engaged_users:,.0f}")

    # Siebel totals
    total_tokens_assigned = 0
    total_tokens_used = 0
    if siebel_devs:
        license_totals = aggregate_license_type_totals(siebel_devs)

        print("\nLicense Type Totals")
        print("-" * 72)
        print(f"{'License Type':<24} {'Assigned':>12} {'Used':>12}")
        for license_type, stats in sorted(license_totals.items()):
            assigned = int(round(stats.get("tokens_assigned", 0)))
            used = int(round(stats.get("tokens_used", 0)))
            print(f"{license_type:<24} {assigned:12,d} {used:12,d}")

        for s in siebel_devs.values():
            try:
                total_tokens_assigned += int(round(s.get("tokens_assigned", 0)))
            except Exception:
                pass
            try:
                total_tokens_used += int(round(s.get("tokens_used", 0)))
            except Exception:
                pass

        total_tokens_remaining = total_tokens_assigned - total_tokens_used
        if total_tokens_remaining < 0:
            total_tokens_remaining = 0

        usage_pct = (
            (total_tokens_used / total_tokens_assigned) * 100
            if total_tokens_assigned
            else 0.0
        )

        print("\nSiebel Totals")
        print("-" * 72)
        print(f"Total tokens assigned : {total_tokens_assigned:,}")
        print(f"Total tokens used     : {total_tokens_used:,}")
        print(f"Total tokens remaining: {total_tokens_remaining:,}")
        print(f"Token usage percent   : {usage_pct:.2f}%")
    else:
        print("\nQuota/balance info was not available in the Copilot metrics payload.")

    print("\nChat Ask Details")
    print("-" * 72)
    if not chat_breakdown:
        print("No chat ask metrics found in response.")
    else:
        for key in sorted(chat_breakdown.keys()):
            print(f"{key:<24}: {chat_breakdown[key]:,.0f}")

    print("\nModel Usage")
    print("-" * 72)
    if not model_usage:
        print("No model-level fields found in response.")
    else:
        preferred_keys = [
            "total_chats",
            "total_chat_requests",
            "total_code_suggestions",
            "total_code_acceptances",
            "total_code_lines_suggested",
            "total_code_lines_accepted",
            "total_engaged_users",
            "total_active_users",
        ]

        for model_name in sorted(model_usage.keys()):
            stats = model_usage[model_name]
            visible_pairs: list[tuple[str, float]] = []

            for key in preferred_keys:
                if key in stats:
                    visible_pairs.append((key, stats[key]))

            # Include any additional numeric model field for transparency.
            for key in sorted(stats.keys()):
                if key not in preferred_keys:
                    visible_pairs.append((key, stats[key]))

            metrics_text = ", ".join(f"{k}={v:,.0f}" for k, v in visible_pairs)
            print(f"{model_name}: {metrics_text}")

    # Siebel developer reporting (license type and token usage)
    siebel_devs = aggregate_siebel_developers(metrics)
    if siebel_devs:
        print("\nSiebel Developer License/Token Report")
        print("-" * 72)
        print(f"{'Developer':<36} {'License':<16} {'Assigned':>10} {'Used':>10}")
        for dev, stats in sorted(siebel_devs.items()):
            license_type = stats.get("license_type", "unknown")
            assigned = int(round(stats.get("tokens_assigned", 0)))
            used = int(round(stats.get("tokens_used", 0)))
            print(f"{dev:<36} {license_type:<16} {assigned:10,d} {used:10,d}")

    # Model token usage / sage cost (per-model totals)
    if token_usage:
        print("\nTotal Copilot Token Usage")
        print("-" * 72)
        grand_total_tokens = 0
        for stats in token_usage.values():
            grand_total_tokens += int(round(stats.get("total_tokens", 0)))
        print(
            f"Grand total tokens (exact if available, otherwise estimated): {grand_total_tokens:,}"
        )
        print("\nPer-model Token Usage (exact counts when available, otherwise estimated)")
        print("-" * 72)
        grand_total_tokens = 0
        grand_total_cost = 0.0
        for model_name, stats in sorted(token_usage.items()):
            inp = int(round(stats.get("input_tokens", 0)))
            out = int(round(stats.get("output_tokens", 0)))
            cached = int(round(stats.get("cached_tokens", 0)))
            total = float(stats.get("total_tokens", 0.0))
            grand_total_tokens += int(round(total))
            rate = 0.0
            if rate_map:
                # try exact then lowercase match
                rate = rate_map.get(model_name, rate_map.get(model_name.lower(), 0.0))
            cost = total * float(rate)
            grand_total_cost += cost
            print(
                f"{model_name:<30} in={inp:10,d} out={out:10,d} cached={cached:10,d} total={int(round(total)):12,d} rate={rate:.8f} cost={cost:,.6f}"
            )

        print("-" * 72)
        print(f"Grand total tokens: {grand_total_tokens:,}")
        print(f"Estimated sage cost: {grand_total_cost:,.6f}")


def main() -> None:
    args = parse_args()

    # Load rates file early so personal mode can estimate costs too
    rate_map = load_rate_map(getattr(args, "rates_file", None))

    # --- Save config if requested ---
    if args.save_config:
        config = load_config()
        if args.org:
            config["org"] = args.org
        if args.enterprise:
            config["enterprise"] = args.enterprise
        if args.token:
            config["token"] = args.token
        if args.author:
            config["author"] = args.author
        save_config(config)
        print(f"Saved settings to {CONFIG_PATH}")

    if args.personal:
        since, until = resolve_personal_window(args.since, args.until)
        sync_stats = {
            "log_files": 0,
            "success_lines": 0,
            "parsed_success": 0,
            "new_added": 0,
            "billing_models": 0,
        }
        current_copilot_sku = None

        # Auto-detect git author if not specified
        if not args.author:
            detected = detect_git_author(args.repo_path)
            if detected:
                args.author = detected
                print(f"Auto-detected git author: {detected}")

        if args.record_ask:
            append_personal_ask_event(args.ask_log_file, args.record_ask, args.ask_model)
            print(f"Recorded ask event in {args.ask_log_file}")

        if not args.no_vscode_sync:
            sync_stats = sync_vscode_asks(args.ask_log_file, since, until, args.vscode_log_dir)
            added = sync_stats.get("new_added", 0)
            if added:
                print(f"Auto-imported {added} new ask(s) from VS Code Copilot Chat logs")

            success_lines = sync_stats.get("success_lines", 0)
            parsed_success = sync_stats.get("parsed_success", 0)
            if success_lines > 0:
                capture_pct = (parsed_success / success_lines) * 100
                print(
                    "VS Code success capture: "
                    f"{parsed_success}/{success_lines} ({capture_pct:.1f}%)"
                )

        current_copilot_sku = detect_copilot_sku(
            find_vscode_copilot_log_files(args.vscode_log_dir),
            since,
            until,
        )

        ask_events = filter_personal_ask_events(
            load_personal_ask_events(args.ask_log_file),
            since,
            until,
        )
        git_stats = get_git_daily_stats(args.repo_path, since, until, args.author)
        rows = build_personal_daily_rows(git_stats, ask_events, since, until)

        # Fetch org license quota if org is specified, or try to auto-detect
        org_license_quota = {}
        detected_org = None
        if args.org or args.enterprise:
            org_license_quota = fetch_org_license_quota(args)
        else:
            # Try to auto-detect org from git remote and fetch quota if token available
            detected_org = detect_github_org_from_git()
            if detected_org and args.token:
                # Create temporary args with detected org for quota fetching
                temp_args = argparse.Namespace(
                    org=detected_org,
                    enterprise=None,
                    team=None,
                    token=args.token,
                    since=args.since,
                    until=args.until,
                )
                org_license_quota = fetch_org_license_quota(temp_args)
                if org_license_quota:
                    print(f"[Auto-detected org: {detected_org}]")

        code_kpi_override = None
        if args.org or args.enterprise:
            code_kpi_override = fetch_code_suggestion_kpis_for_scope(args)
        elif detected_org and args.token:
            temp_kpi_args = argparse.Namespace(
                org=detected_org,
                enterprise=None,
                team=None,
                token=args.token,
                since=args.since,
                until=args.until,
            )
            code_kpi_override = fetch_code_suggestion_kpis_for_scope(temp_kpi_args)

        personal_summary = build_personal_summary(
            rows,
            ask_events,
            code_kpi_override=code_kpi_override,
        )
        
        print_personal_summary(
            rows,
            ask_events,
            org_license_quota=org_license_quota,
            current_copilot_sku=current_copilot_sku,
            code_kpi_override=code_kpi_override,
        )

        # Also print structured result block matching the requested format
        # Capture values available from earlier steps
        detected_author = args.author or "unknown"
        auto_imported = 0
        capture_rate_text = ""
        try:
            auto_imported = int(sync_stats.get("new_added", 0))
            success_lines = int(sync_stats.get("success_lines", 0))
            parsed_success = int(sync_stats.get("parsed_success", 0))
            if success_lines:
                capture_rate_text = f"{parsed_success}/{success_lines} ({(parsed_success/success_lines*100):.1f}%)"
            else:
                capture_rate_text = "0/0 (0.0%)"
        except Exception:
            capture_rate_text = "unknown"

        # Build ask model breakdown string
        ask_models = summarize_ask_models(ask_events)

        print("\n### Result\n")
        print(f"- Auto-detected git author: `{detected_author}`")
        print(f"- Auto-imported `{auto_imported}` new ask(s) from VS Code Copilot Chat logs")
        print(f"- Capture rate: `{capture_rate_text}`\n")

        print("### Summary\n")
        if current_copilot_sku:
            print(f"- Copilot SKU: `{current_copilot_sku}`")
        personal_license_types = personal_summary.get("license_types", ["unknown"])
        print(f"- License type(s): `{', '.join(personal_license_types)}`")
        print(f"- Days covered: `{len(rows)}`")
        print(f"- Date range: `{rows[0]['date']} to {rows[-1]['date']}`")
        print(f"- Commits: `{sum(r['commits'] for r in rows)}`")
        print(f"- Files changed: `{sum(r['files_changed'] for r in rows)}`")
        print(f"- Lines added: `{sum(r['lines_added'] for r in rows):,}`")
        print(f"- Lines deleted: `{sum(r['lines_deleted'] for r in rows)}`")
        print(f"- Net lines: `{sum(r['net_lines'] for r in rows):,}`")
        print(f"- Asks logged: `{len(ask_events)}`")
        print(f"- Ask characters: `{sum(e.get('ask_length', 0) for e in ask_events)}`")
        kpis = personal_summary.get("code_suggestion_kpis", {})
        if kpis.get("available"):
            print(f"- Shown: `{int(kpis.get('shown', 0)):,}`")
            print(f"- Accepted: `{int(kpis.get('accepted', 0)):,}`")
            print(f"- Lines of code accepted: `{int(kpis.get('lines_of_code_accepted', 0)):,}`")
            print(f"- % Acceptance: `{float(kpis.get('acceptance_pct', 0.0)):.2f}%`")
        else:
            print("- Shown: `N/A`")
            print("- Accepted: `N/A`")
            print("- Lines of code accepted: `N/A`")
            print("- % Acceptance: `N/A`")
        transcript_data = personal_summary.get("transcript_tokens", {})
        if transcript_data.get("total_tokens", 0) > 0:
            print(f"- Total tokens consumed: `{transcript_data['total_tokens']:,}`")
            print(f"- Total input tokens: `{transcript_data['input_tokens']:,}`")
            print(f"- Total output tokens: `{transcript_data['output_tokens']:,}`")
        quota_info = extract_personal_token_quota(ask_events)
        if quota_info.get("total_assigned_tokens", 0) > 0 or quota_info.get("total_used_tokens", 0) > 0:
            print(f"- Total tokens assigned: `{quota_info['total_assigned_tokens']:,}`")
            print(f"- Total tokens used: `{quota_info['total_used_tokens']:,}`")
            print(f"- Total tokens remaining: `{quota_info['total_remaining_tokens']:,}`")
            print(f"- Token usage percent: `{quota_info['usage_percent']:.2f}%`")
        billing_summary = personal_summary.get("billing_summary", {})
        if billing_summary.get("total_tokens", 0) or billing_summary.get("estimated_cost", 0.0):
            print(f"- Ask input tokens: `{int(billing_summary.get('input_tokens', 0)):,}`")
            print(f"- Ask output tokens: `{int(billing_summary.get('output_tokens', 0)):,}`")
            print(f"- Ask total tokens: `{int(billing_summary.get('total_tokens', 0)):,}`")
            print(f"- Estimated ask cost: `{float(billing_summary.get('estimated_cost', 0.0)):,.6f}`")
        # model breakdown lines
        for model, count in sorted(ask_models.items()):
            print(f"- Ask model breakdown:`{model}`: `{count}`")

        # Estimate tokens and cost from personal asks (heuristic)
        token_est = estimate_personal_tokens(ask_events, rate_map=rate_map)
        if token_est.get("total_tokens", 0):
            print("\nEstimated Personal Token Usage")
            print("-" * 72)
            print(f"Asks counted      : {token_est['total_asks']}")
            print(f"Total characters   : {token_est['total_chars']}")
            print(f"Estimated tokens   : {token_est['total_tokens']:,}")
            if rate_map:
                print(f"Estimated sage cost: {token_est['estimated_cost']:,.6f}")
            if token_est.get("per_model_tokens"):
                print("Per-model estimated tokens:")
                for m, t in sorted(token_est["per_model_tokens"].items()):
                    print(f"  {m:<20}: {t:,}")

        if args.csv_out:
            write_personal_csv(rows, args.csv_out)
            print(f"\nWrote personal CSV summary: {args.csv_out}")

        if args.trend_out:
            write_personal_trend_report(rows, args.trend_out)
            print(f"Wrote personal trend report: {args.trend_out}")

        if args.excel_out:
            write_personal_excel(rows, ask_events, args.excel_out)
            print(f"Wrote personal Excel summary: {args.excel_out}")
        return

    metrics = fetch_metrics(args)
    rate_map = load_rate_map(getattr(args, "rates_file", None))
    print_summary(metrics, rate_map=rate_map)

    # Optional Siebel developer CSV output
    if getattr(args, "siebel_report", None):
        siebel_devs = aggregate_siebel_developers(metrics)
        if siebel_devs:
            write_siebel_csv(siebel_devs, args.siebel_report)
            print(f"\nWrote Siebel developer CSV: {args.siebel_report}")
        else:
            print("\nNo Siebel developer fields found to write report.")

    if getattr(args, "license_summary_report", None):
        siebel_devs = aggregate_siebel_developers(metrics)
        license_totals = aggregate_license_type_totals(siebel_devs)
        if license_totals:
            write_license_summary_csv(license_totals, args.license_summary_report)
            print(f"Wrote license summary CSV: {args.license_summary_report}")
        else:
            print("No license-type token fields found to write summary report.")

    # Optional per-model cost CSV
    if getattr(args, "cost_report", None):
        token_usage = aggregate_model_token_usage(metrics)
        write_cost_csv(token_usage, args.cost_report, rate_map)
        print(f"\nWrote per-model cost CSV: {args.cost_report}")

    daily_rows = build_daily_rows(metrics)

    if args.csv_out:
        write_daily_csv(daily_rows, args.csv_out)
        print(f"\nWrote CSV summary: {args.csv_out}")

    if args.trend_out:
        write_trend_report(daily_rows, args.trend_out)
        print(f"Wrote trend report: {args.trend_out}")

    if args.excel_out:
        siebel_devs = aggregate_siebel_developers(metrics)
        write_org_excel(daily_rows, siebel_devs, metrics, rate_map, args.excel_out)
        print(f"Wrote Excel summary: {args.excel_out}")

    if args.raw_json:
        print("\nRaw JSON payload")
        print("-" * 72)
        print(json.dumps(metrics, indent=2))


if __name__ == "__main__":
    main()
